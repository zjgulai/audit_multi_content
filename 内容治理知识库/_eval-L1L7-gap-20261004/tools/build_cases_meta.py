#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""组装某件 skill 的 `cases-meta.json`（A 卷 / 调用题 / B 卷）——**复用既有用例定义，不改任何用例**。

支持三种既有产物形态（自动按内容判别，不靠文件名猜）：
  ① `{results:[{id,type,expected,tags,title,priority}]}` + 单独的 prompts 文件
     prompts 形态：`[{id,n,prompt}]`
  ② `{cases:[{id,type,expected,tags,title,prompt,priority}]}` —— 自带 prompt
  ③ `.xlsx` 用例表：表头含 序号/类型/标题/提示词/预期结果/标签/优先级

输出（与 L0 的 cases-meta.json 同构，judge.py / 判分脚本可直接吃）：
  {"skill":..., "a_cases":[...], "call_cases":[...], "b_cases":[...]}
  `type == "调用"` 进 call_cases；其余进 a_cases。

★ 只搬运，不推断、不改写。缺 `prio` 填 "P1"（既有产物默认）。
"""
import argparse
import collections
import json
import os
import re


def _norm_case(cid, typ, title, prompt, expected, tags, prio):
    n = int(cid) if str(cid).isdigit() else str(cid)
    return {"n": n, "title": title or "", "type": typ, "prompt": prompt or "",
            "expected": expected, "tags": tags or "", "prio": prio or "P1"}


# 中文键 → 英文键（既有产物两种写法都有：kb-recall 用中文键，media 用英文键）
CN = {"类型": "type", "标题": "title", "提示词": "prompt", "预期结果": "expected",
      "标签": "tags", "优先级": "priority", "序号": "id", "编号": "id"}


def _en(d):
    """把一条记录的键归一成英文（不改值）。"""
    if not isinstance(d, dict):
        return d
    out = dict(d)
    for cn, en in CN.items():
        if cn in out and en not in out:
            out[en] = out[cn]
    return out


def load_side(path, prompts=None):
    """返回 (cases, warnings)。path 可以是 json 或 xlsx。"""
    if not path:
        return [], []
    warn = []
    if path.endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(path)
        ws = wb[wb.sheetnames[0]]
        hdr = [str(c.value).strip() if c.value else "" for c in ws[1]]
        col = {h: i for i, h in enumerate(hdr)}

        def g(row, name):
            i = col.get(name)
            return row[i].value if i is not None and i < len(row) else None

        out = []
        for row in ws.iter_rows(min_row=2):
            if g(row, "序号") is None:
                continue
            out.append(_norm_case(g(row, "序号"), g(row, "类型"), g(row, "标题"),
                                  g(row, "提示词"), g(row, "预期结果"),
                                  g(row, "标签"), g(row, "优先级")))
        return out, warn

    d = json.load(open(path, encoding="utf-8"))
    if isinstance(d, dict) and "cases" in d:            # 形态②（{cases:[...]}）
        out = []
        for c0 in d["cases"]:
            c = _en(c0)
            out.append(_norm_case(c.get("id"), c["type"], c.get("title"), c.get("prompt"),
                                  c["expected"], c.get("tags"),
                                  c.get("priority") or c.get("prio")))
        return out, warn
    labels = d["results"] if isinstance(d, dict) else d  # 形态①（含裸 list，可能是中文键）
    labels = [_en(x) for x in labels]
    pmap = {}
    if prompts:
        pd = json.load(open(prompts, encoding="utf-8"))
        for x in pd:
            pmap[str(x.get("id"))] = x.get("prompt")
            if "n" in x:
                pmap[str(x["n"])] = x.get("prompt")
    out = []
    for r in labels:
        cid = str(r["id"])
        # 形态①-b：记录**自带** prompt（kb-recall 的 cases-a.json 就是裸 list 且自带「提示词」）
        p = r.get("prompt")
        if p is None:
            p = pmap.get(cid)
            if p is None:
                warn.append(f"取不到 prompt: id={cid}")
        out.append(_norm_case(cid, r["type"], r.get("title"), p, r["expected"],
                              r.get("tags"), r.get("priority") or r.get("prio")))
    return out, warn


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--skill", required=True)
    ap.add_argument("--a", required=True, help="A 卷用例定义（json 或 xlsx）")
    ap.add_argument("--a-prompts", default=None, help="形态①才需要")
    ap.add_argument("--b", default=None)
    ap.add_argument("--b-prompts", default=None)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()

    a_all, w1 = load_side(a.a, a.a_prompts)
    b_all, w2 = load_side(a.b, a.b_prompts)
    a_cases = [c for c in a_all if c["type"] != "调用"]
    call_cases = [c for c in a_all if c["type"] == "调用"]

    out = {"skill": a.skill,
           "a_cases": sorted(a_cases, key=lambda c: (isinstance(c["n"], str), c["n"])),
           "call_cases": sorted(call_cases, key=lambda c: (isinstance(c["n"], str), c["n"])),
           "b_cases": sorted(b_all, key=lambda c: (isinstance(c["n"], str), c["n"]))}
    json.dump(out, open(a.out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    print(f"{a.skill}: a_cases={len(a_cases)} call_cases={len(call_cases)} b_cases={len(b_all)}")
    print("  A 卷类型:", dict(collections.Counter(c["type"] for c in a_cases)))
    if b_all:
        print("  B 卷类型:", dict(collections.Counter(c["type"] for c in b_all)))
    print("  tags 非空:", sum(1 for c in a_cases + call_cases if c["tags"]),
          "｜空 prompt:", sum(1 for c in a_cases + call_cases if not c["prompt"]))
    for w in w1 + w2:
        print("  ⚠️", w)
    return 0 if not (w1 or w2) else 1


if __name__ == "__main__":
    raise SystemExit(main())
