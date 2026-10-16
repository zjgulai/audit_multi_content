#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""回填员脚本：把 3 轮盲测会话逐字回填成 results.json / score-input.json / Excel。

隔离口径（本轮实测必需）：只有 parentSession == MY_SESSION 的会话才算本轮盲测。
理由见 _eval/CONCURRENCY-NOTE.md —— 用例 33 的被测会话自行跑了一轮完整评测，
向本目录 spawn 了 161 个同题面会话（parentSession=92c12bd1…），必须按 parent 隔离。

判定口径（root-skills-eval v3.19.1 Step 4）：
  触发 = 该会话用 skill 工具加载了 kb-check（或注入 <skill_content name="kb-check">）
  追问 = 未加载 kb-check 且最终回复以追问澄清收尾
  不触发 = 其余
  调用 = 加载 / 未加载（主语是本 skill）
"""
from __future__ import annotations

import json
import os
import sys
from datetime import datetime

import zstandard

HERE = os.path.dirname(os.path.abspath(__file__))
EV = os.path.dirname(HERE)
MY_SESSION = "f41c4909-8674-451f-b32b-1358305ce38e"   # 评测主会话（回填员）
TARGET = "kb-check"
SESS = os.path.expanduser("~/.dsh/sessions")
NS_KEY = "~5185~5BB9~6CBB~7406"
dec = zstandard.ZstdDecompressor()
sys.path.insert(0, HERE)
import session_evidence as se  # noqa: E402


def load_cases(name):
    return json.load(open(os.path.join(EV, name), encoding="utf-8"))


def collect(cases, out_name):
    want = {str(c["提示词"]).strip(): k for k, c in cases.items()}
    hits = {}
    for aid, fp in se._find_all_session_files():
        try:
            a = se.analyse(fp)
        except Exception:
            continue
        if a.get("parent_session") != MY_SESSION:
            continue
        for ut in a["user_texts"]:
            if ut in want:
                cid = want[ut]
                a["agent_id"], a["case_id"] = aid, cid
                hits.setdefault(cid, []).append(a)
                break
    json.dump(hits, open(os.path.join(HERE, out_name), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    return hits


def verdict(case, a):
    typ = case["类型"]
    if typ == "调用":
        return "加载" if a["loaded_target_skill"] else "未加载"
    if a["loaded_target_skill"]:
        return "触发"
    return "追问" if a["ends_with_question"] else "不触发"


def build(cases, hits, tag):
    results, score_input, ledger, bad = {}, [], {}, []
    for cid in cases:
        case = cases[cid]
        rows = sorted(hits.get(cid, []), key=lambda a: a.get("createdAt") or 0)
        if len(rows) != 3:
            bad.append((cid, len(rows)))
        rounds = []
        for a in rows[:3]:
            rounds.append({"结论": verdict(case, a), "回复": a["reply"],
                           "_agent": a["agent_id"], "_skills": a["skill_tool_calls"] + a["injected_skills"],
                           "_ts": datetime.fromtimestamp((a.get("createdAt") or 0) / 1000).strftime("%Y-%m-%d %H:%M:%S")})
        while len(rounds) < 3:
            rounds.append({"结论": "", "回复": "", "_agent": "", "_skills": [], "_ts": ""})
        results[cid] = {"标题": case["标题"], "类型": case["类型"], "提示词": case["提示词"],
                        "预期结果": case["预期结果"], "标签": case["标签"], "优先级": case["优先级"],
                        "轮次": rounds}
        score_input.append({"id": cid, "type": case["类型"], "expected": case["预期结果"],
                            "tags": case["标签"], "title": case["标题"],
                            "rounds": [r["结论"] for r in rounds]})
        ledger[cid] = [{"round": i + 1, "agent_id": r["_agent"], "verdict": r["结论"],
                        "skills": r["_skills"], "ts": r["_ts"], "reply_chars": len(r["回复"])}
                       for i, r in enumerate(rounds)]
    json.dump(results, open(os.path.join(EV, f"results-{tag}.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    json.dump({"skill_name": TARGET, "results": score_input},
              open(os.path.join(HERE, f"score-input-{tag}.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    json.dump(ledger, open(os.path.join(HERE, f"run-ledger-{tag}.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    return results, bad


def fill_excel(path, results):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["Sheet1"]
    idx = {str(ws.cell(r, 1).value).strip(): r for r in range(2, ws.max_row + 1)}
    for cid, c in results.items():
        r = idx.get(str(cid))
        if not r:
            continue
        for i, rd in enumerate(c["轮次"]):
            ws.cell(r, 8 + i).value = f"{rd['结论']}｜{rd['回复']}"
        ws.cell(r, 11).value = c["轮次"][0]["_ts"][:10]
    wb.save(path)


if __name__ == "__main__":
    A = load_cases("cases-a.json")
    hits = collect(A, "evidence-mine-a.json")
    res, bad = build(A, hits, "a")
    fill_excel(os.path.join(EV, "kb-check.xlsx"), res)
    print(f"A 卷 {len(A)} 条；三轮齐备异常：{bad or '无'}")
    from collections import Counter
    c = Counter(r["结论"] for v in res.values() for r in v["轮次"])
    print("结论分布：", dict(c))
    for i in range(3):
        cc = Counter(v["轮次"][i]["结论"] for v in res.values() if v["类型"] != "调用")
        print(f"  R{i+1}:", dict(cc))
