#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
《内容治理知识库搭建任务详情.xlsx》(定稿版 v14)

按 PM 反馈返工：
  · 只保留「任务详情」一页，其余 sheet 全部删除
  · A1 写标题：内容治理Agent排期
  · 日期第一列，一行一周（写该周的时间段），不写星期
表头：日期 / 任务名称 / H2目标 / 预期交付物 / 负责人 / 当前状态
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DESKTOP = Path("/Users/lute/Desktop")
SRC = DESKTOP / "内容治理知识库搭建任务详情.xlsx"

TITLE = "内容治理Agent排期"
OWNER = "@Pray-周健 @刘华丽"

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=15, color="1F3864")
NAME_FONT = Font(bold=True, size=11, color="1F3864")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP_WRAP = Alignment(vertical="top", wrap_text=True)
CTR = Alignment(vertical="top", horizontal="center", wrap_text=True)
BAND = PatternFill("solid", fgColor="F2F5FA")
MARK = PatternFill("solid", fgColor="FFF2CC")
GRAY = PatternFill("solid", fgColor="EDEDED")

ROWS = [
["9/28–9/30", "内容治理知识库方案 + 守衡接通",
 "输出《内容治理知识库方案》；接通飞书守衡，让守衡能只读调用 DSH 中的知识库；同时定完外部标准源白名单、指南 PDF 授权、规则库表处置，并派出权限申请",
 "《内容治理知识库方案》；飞书守衡只读调用 DSH 知识库的接通验证（健康检查、P0 匹配、检索端点可用）；外部标准源白名单；权限申请单",
 OWNER, "待启动"],

["10/1–10/7", "国庆放假", "假期不排节点", "—", "—", "—"],

["10/8–10/10", "知识库底座修复",
 "装载位归位、L1 索引两缺陷修复、L6 重建、真源统一（条款号、锚点、source_sha256 只在权威库生成）",
 "装载位；L6 索引与自测结果",
 OWNER, "待启动"],

["10/12–10/16", "文字链路首次跑通",
 "回填 owner、effective_from、risk_level 与 modality 字段；清掉三个假绿默认值；写入时最小脱敏；9 场景覆盖对账与蒸馏工单，启用 W-4 段；文字链路跑通",
 "基线表（健康分、规则孤岛率、检索命中）；9 场景覆盖表；蒸馏工单；W-4 编号规则 ≥1 条；10 条真实内容跑通记录",
 OWNER, "待启动"],

["10/19–10/23", "外部条款入库 + 图片 OCR 起步",
 "抓自有站 166 条、平台政策 22 条、FDA 24 条、FTC 并逐字回验，编 E-2 发条款号；新建 `kb-clause-extractor`；图片 OCR 抽字跑通，以系统 OCR 为准、带指纹可回源",
 "E-2 条款与条款号台账；逐字回验报告；`kb-clause-extractor`；图片 OCR 抽取结果",
 OWNER, "待启动"],

["10/26–10/30", "权威指南条款 + 视频关键帧",
 "WHO、AAP 指南条款入库；视频抽样与关键帧提取跑通，用 10–20 条小样试判",
 "WHO、AAP 条款号；关键帧与视频试判记录",
 OWNER, "待启动"],

["11/2–11/6", "红线硬拦截 + 图片链路 + 输出规格",
 "7 类红线确定性硬拦截（命中即拦、不靠 AI）并由法务逐条确认；补中英同义、英文屈折、连字符变体与德英同义族；图片走通完整审查链路；定输出规格、结论路由、处理时限、缺料处置、三态与「查不到」出口、置信度分流、被审对象标定",
 "7 类红线规则集；同义与变体表；图片三态结论与依据锚点；输出规格与路由、时限规则",
 OWNER, "待启动"],

["11/9–11/13", "样本清洗 + 视频 ASR",
 "208 条内部样本清洗与分级，先蒸 40 条（红线 6 条优先）；视频字幕、旁白 ASR 与时间码锚点跑通",
 "干净样本集与分级表；样本规则（带条款号）；视频字幕与旁白抽取结果",
 OWNER, "待启动"],

["11/16–11/20", "产品事实库 + 场景规则 + 视频出结论",
 "产品事实库与三级审批队列上线；九场景补规则（画面类优先）；视频出结论并定画面类结论上限（只判需复核、不单独驳回）；飞书群接入；新建 `brand-product-fact-curator`、`kb-ingest-extractor`",
 "产品事实库；九场景规则清单；视频结论与结论上限规则；飞书群接入",
 OWNER, "待启动"],

["11/23–11/27", "四项能力验收",
 "扫描件 PDF 转图认字；批量（单次 ≤50 个文件、并发 100–500 条）；六项速度实测；打通回流通道；端到端跑通；文字 MVP、图片、视频、批量与回流验收",
 "四项验收报告；端到端跑通记录；速度实测结果；本期不做项清单",
 OWNER, "待启动"],

["11/30–12/4", "准确率评测",
 "用 100 条标准答案（文字 40、图片 30、视频 30）做双盲评测；补对抗题、扰动题、德语题；完成五类归因；B 组一律拒答",
 "标准答案与标注；五类归因表；对抗题、扰动题、德语题",
 OWNER, "待启动"],

["12/7–12/11", "第一轮修复",
 "集中修影响最大的一类问题；改前留可回退版本；知识库数据问题重跑各层、重建索引、补字段；准确率问题补规则、调门限、补样本；skill 问题在副本上修",
 "修复记录；回归重测报告；可回退版本",
 OWNER, "待启动"],

["12/14–12/18", "指标结算",
 "结算五项指标（三态准确率 ≥90%、驳回精确率 ≥95%、通过误放率 ≤5%、需复核召回率 ≥95%、红线漏放率 0）与人工负担降幅 ≥50%；未达标项当天定处置",
 "指标结算表；处置决定；规则命中率与回流进库条数",
 OWNER, "待启动"],

["12/21–12/24", "锁版发版",
 "5% 分层抽查；发四类库 kb_version 与整体 release_id、写 versions.jsonl；条款号核验；半衰期按域取值；源侧零改动证明；脱敏复核；新 skill 跑 `root-skills-eval`；跑 `acceptance_disjoint_check` 自检",
 "抽查记录；版本台账；条款号核验结果；脱敏复核记录；skill 评测报告；交接文档与 v1.1 清单",
 OWNER, "待启动"],

["12/25–12/31", "上线判定",
 "全量回归；缺陷修复（只改副本、不动原版）；复测；出正式评估报告；上线判定；核对输入规格（图片 ≤20MB、视频 ≤500MB 且 ≤10 分钟、单任务 ≤50 个文件）与审查边界",
 "正式评估报告；上线判定结论（或《未达标说明》）",
 OWNER, "待启动"],
]

wb = Workbook()
ws = wb.active
ws.title = "任务详情"

# 标题
ws["A1"] = TITLE
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
ws["A1"].font = TITLE_FONT
ws["A1"].alignment = Alignment(vertical="center", horizontal="left")
ws.row_dimensions[1].height = 30

# 表头
ws.append(["日期", "任务名称", "H2目标", "预期交付物", "负责人", "当前状态"])
for row in ROWS:
    ws.append(row)

WIDTHS = [12, 26, 56, 50, 20, 10]
for i, w in enumerate(WIDTHS, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

for c in range(1, 7):
    cell = ws.cell(row=2, column=c)
    cell.fill = HDR_FILL
    cell.font = HDR_FONT
    cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    cell.border = BORDER
ws.row_dimensions[2].height = 26

for r in range(3, ws.max_row + 1):
    for c in range(1, 7):
        cell = ws.cell(row=r, column=c)
        cell.alignment = CTR if c in (1, 5, 6) else TOP_WRAP
        cell.border = BORDER
        if r % 2 == 0 and cell.fill.fgColor.rgb in (None, "00000000"):
            cell.fill = BAND
    ws.cell(row=r, column=2).font = NAME_FONT
    nm = str(ws.cell(row=r, column=2).value)
    if nm == "国庆放假":
        for c in range(1, 7):
            ws.cell(row=r, column=c).fill = GRAY
    elif nm in ("文字链路首次跑通", "外部条款入库 + 图片 OCR 起步", "权威指南条款 + 视频关键帧",
                "四项能力验收"):
        ws.cell(row=r, column=2).fill = MARK
        ws.cell(row=r, column=1).fill = MARK

ws.freeze_panes = "C3"
ws.auto_filter.ref = f"A2:F{ws.max_row}"

wb.save(SRC)
print("updated ->", SRC)
print("sheets  ->", wb.sheetnames)
print("rows    ->", ws.max_row - 2, "个节点")
