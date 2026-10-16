#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把每条回复里「决定归口」的句子抽出来，供回填员逐条复读定案。

理由：知识库型 skill 的盲测回复里一定会**提到**本 skill 的名字（哪怕是在解释
「这条不归我」）。用「提到引擎名/路径」当触发判据会大面积误判，必须复读定案。

用法： /usr/bin/python3 extract-verdicts.py [--round N] [--cases 1,2,3] > verdicts.txt
"""
import json
import os
import re
import sys
import glob
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
EVAL = os.path.dirname(HERE)
NAME = "kb-media-extractor"

OTHER = (r"kb-l0-feishu|kb-ingest-extractor|kb-clause-extractor|kb-pipeline|kb-check|kb-repair|"
         r"kb-recall|kb-canonical-lookup|kb-p0-check|kb-manage|图片安全审查|视频安全审查|文案合规审查|"
         r"内容审查总控|shuorenhua")
ROUTE = re.compile(r"(归|走|交给|应由|属于|让位|路由到|转给|是)\s*[`\*\s]*(" + OTHER + r")")
NEG = re.compile(r"(不|非|不是|不该由|不属于|不在|不接|不碰|不管|不适用|没有)\s*[^。；\n]{0,20}" + NAME +
                 r"|" + NAME + r"[^。；\n]{0,25}(不是|不接|不碰|不管|不适用|不干|之外|范围外)")
NOTRIG = re.compile(r"不触发|不该触发|不作为承接|我不接")


def sentences(t):
    t = re.sub(r"\n+", "\n", t)
    parts = re.split(r"(?<=[。！？\n])", t)
    return [p.strip() for p in parts if p.strip()]


def main():
    rounds = [int(a) for a in sys.argv[1:] if a.isdigit()] or [1, 2, 3]
    only = None
    for a in sys.argv[1:]:
        if "," in a:
            only = {int(x) for x in a.split(",")}
    rows = {}
    for f in sorted(glob.glob(os.path.join(HERE, "a-round-*.jsonl"))):
        rnd = int(os.path.basename(f).split("-")[-1].split(".")[0])
        if rnd not in rounds:
            continue
        for line in open(f, encoding="utf-8"):
            if not line.strip():
                continue
            r = json.loads(line)
            rows.setdefault(r["n"], {})[rnd] = r
    ws = openpyxl.load_workbook(os.path.join(EVAL, f"{NAME}.xlsx"))["触发评测用例"]
    hdr = [c.value for c in ws[1]]
    exp = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        d = dict(zip(hdr, r))
        exp[d["序号"]] = d
    for n in sorted(rows):
        if only and n not in only:
            continue
        e = exp[n]
        print("=" * 110)
        print(f"#{n} [{e['类型']}] {e['标题']}")
        print(f"  预期: {e['预期结果']}")
        print(f"  提示词: {e['提示词']}")
        for rnd in rounds:
            r = rows[n].get(rnd)
            if not r:
                continue
            t = r["reply"]
            head = t[:260].replace("\n", " ⏎ ")
            hits, routes, negs = [], [], []
            for s in sentences(t):
                if NAME in s and len(s) < 400:
                    hits.append(s.replace("\n", " "))
                if ROUTE.search(s):
                    routes.append(s.replace("\n", " ")[:220])
                if NEG.search(s) or NOTRIG.search(s):
                    negs.append(s.replace("\n", " ")[:220])
            print(f"  ── 第 {rnd} 轮 ── oracle={r['oracle']['verdict']} strong={len(r['oracle']['strong'])} med={len(r['oracle']['medium'])}")
            print(f"     [开头] {head}")
            for s in negs[:4]:
                print(f"     [否定/不触发] {s}")
            for s in routes[:3]:
                print(f"     [路由给谁] {s}")
            for s in hits[:3]:
                print(f"     [提到本 skill] {s[:220]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
