#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 3 轮盲测回填成 results.json / score-input.json / Excel 回填。

结论口径（写进评测过程，三轮一致）：
  触发   = 该会话加载了 kb-canonical-lookup（skill 工具点本名）
  不触发 = 未加载本 skill —— 含加载了别的 skill、未加载任何 skill
  追问   = 未加载任何 skill，且回复以追问澄清为主（边界情况专用）
  加载/未加载 = 调用类（主语是本 skill）
"""
import argparse
import json
import os

WORK = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(WORK)
NAME = "kb-canonical-lookup"
TARGET = "kb-canonical-lookup"

OVERRIDES = {}
ov_path = os.path.join(WORK, "overrides.json")
if os.path.exists(ov_path):
    OVERRIDES = json.load(open(ov_path, encoding="utf-8"))


def conclusion(rec):
    cid = str(rec["id"])
    if cid in OVERRIDES:
        return OVERRIDES[cid]
    if TARGET in (rec.get("skill_calls") or []):
        return "加载" if rec.get("type") == "调用" else "触发"
    if rec.get("type") == "调用":
        return "未加载"
    return rec.get("auto_conclusion") or "不触发"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--rounds", nargs=3, required=True, help="round1.json round2.json round3.json")
    ap.add_argument("--outdir", default=WORK)
    ap.add_argument("--date", default="2026-09-30")
    args = ap.parse_args()

    cases = json.load(open(os.path.join(WORK, "cases_a.json"), encoding="utf-8"))["cases"]
    order = [str(c["id"]) for c in cases]
    meta = {str(c["id"]): c for c in cases}

    rounds = [json.load(open(p, encoding="utf-8")) for p in args.rounds]

    results = {}
    score_results = []
    for cid in order:
        rounds_data = []
        for r in rounds:
            rec = r.get(cid)
            if rec is None:
                rounds_data.append({"结论": "", "回复": "", "session": ""})
                continue
            rounds_data.append({
                "结论": conclusion(rec),
                "回复": rec.get("reply", ""),
                "session": rec.get("session", ""),
            })
        c = meta[cid]
        results[cid] = {
            "标题": c["title"], "类型": c["type"], "提示词": c["prompt"],
            "预期结果": c["expect"], "标签": c["tags"], "优先级": c["prio"],
            "轮次": rounds_data,
        }
        score_results.append({
            "id": cid, "type": c["type"], "expected": c["expect"],
            "tags": c["tags"], "title": c["title"],
            "rounds": [x["结论"] for x in rounds_data],
        })

    with open(os.path.join(args.outdir, "results.json"), "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    with open(os.path.join(args.outdir, "score-input.json"), "w", encoding="utf-8") as f:
        json.dump({"skill_name": NAME, "results": score_results}, f, ensure_ascii=False, indent=2)

    # Excel 回填
    from openpyxl import load_workbook
    xp = os.path.join(ROOT, f"{NAME}.xlsx")
    wb = load_workbook(xp)
    ws = wb["Sheet1"]
    head = [c.value for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(head)}
    for row in ws.iter_rows(min_row=2):
        cid = str(row[idx["序号"] - 1].value)
        if cid not in results:
            continue
        for i, col in enumerate(("1轮执行结果", "2轮执行结果", "3轮执行结果")):
            ws.cell(row=row[0].row, column=idx[col]).value = results[cid]["轮次"][i]["结论"]
        ws.cell(row=row[0].row, column=idx["执行日期"]).value = args.date
    wb.save(xp)

    print(f"[ok] results.json / score-input.json / {NAME}.xlsx 回填完成（{len(order)} 条 × 3 轮）")
    empty = [cid for cid, v in results.items() if any(not x["回复"] for x in v["轮次"])]
    print("回复为空的用例:", empty)


if __name__ == "__main__":
    main()
