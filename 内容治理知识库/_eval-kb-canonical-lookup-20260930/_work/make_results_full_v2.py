#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 v1.3.0 的 results.json（供 build-testcases-html.py 用）+ 回填 Excel。

结论口径 = 口径 A（与 v1.2.0 发布的数字同尺）。
回复 = 被测会话的真实原文逐字；以 ask_user_question 收尾的会话，其提问正文由
collect_v2.py 逐字回填为回复（不截断、不摘要）。
"""
import argparse
import json
import os

WORK = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(WORK)
NAME = "kb-canonical-lookup"
TARGET = "kb-canonical-lookup"


def conclusion(rec):
    if rec.get("type") == "调用":
        return "加载" if TARGET in (rec.get("skill_calls") or []) else "未加载"
    return "触发" if TARGET in (rec.get("skill_calls") or []) else "不触发"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--rounds", nargs=3, required=True)
    ap.add_argument("--cases", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--xlsx", default="")
    ap.add_argument("--date", default="2026-09-30")
    args = ap.parse_args()

    cases = json.load(open(args.cases, encoding="utf-8"))["cases"]
    order = [str(c["id"]) for c in cases]
    meta = {str(c["id"]): c for c in cases}
    rounds = [json.load(open(p, encoding="utf-8")) for p in args.rounds]

    results = {}
    for cid in order:
        cells = []
        for r in rounds:
            rec = r.get(cid)
            cells.append({"结论": conclusion(rec) if rec else "",
                          "回复": (rec or {}).get("reply", ""),
                          "session": (rec or {}).get("session", "")})
        c = meta[cid]
        results[cid] = {"标题": c["title"], "类型": c["type"], "提示词": c["prompt"],
                        "预期结果": c["expect"], "标签": c["tags"], "优先级": c["prio"],
                        "轮次": cells}

    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)

    empty = [cid for cid, v in results.items() if any(not x["回复"] for x in v["轮次"])]
    print(f"[ok] {args.out}（{len(order)} 条 × 3 轮）；回复为空: {empty}")

    if args.xlsx:
        from openpyxl import load_workbook
        wb = load_workbook(args.xlsx)
        ws = wb["Sheet1"]
        head = [c.value for c in ws[1]]
        idx = {h: i + 1 for i, h in enumerate(head)}
        for row in ws.iter_rows(min_row=2):
            cid = str(row[idx["序号"] - 1].value)
            if cid not in results:
                continue
            for i, col in enumerate(("1轮执行结果", "2轮执行结果", "3轮执行结果")):
                ws.cell(row=row[0].row, column=idx[col]).value = results[cid]["轮次"][i]["结论"]
            ws.cell(row=row[0].row, column=idx["执行日期"]).value = args.date
        wb.save(args.xlsx)
        print("[ok] 回填", args.xlsx)


if __name__ == "__main__":
    main()
