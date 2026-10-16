#!/usr/bin/env python3
"""
split-gate.py - 留出验证门（held-out validation gate）

借鉴 microsoft/SkillOpt 的验证门控设计：修复候选由「引导集」(guide) 生成，
接受与否只看「留出集」(holdout) 的表现，避免修复过拟合到同一批触发用例。
配合 --baseline 实现 gate_no_regression（留出集只升不降）语义。

输入是触发评测结果（score-trigger-eval.py 的输入格式，JSON 或 Excel）：
    {
      "skill_name": "example-skill",
      "results": [
        {"id": 1, "type": "正例", "expected": "应触发 example-skill",
         "rounds": ["触发", "触发", "不触发"]},
        ...
      ]
    }

用法:
    # 查看 guide/holdout 切分与各自 3 轮指标
    python3 scripts/split-gate.py results.json --seed 42 --holdout-ratio 0.2

    # 生成修复前基线（写 gate-baseline.json，供修复轮对比）
    python3 scripts/split-gate.py results.json --seed 42 --baseline-out gate-baseline.json

    # 与基线对比，判定留出集是否回归（--baseline 输入）
    python3 scripts/split-gate.py results.json --baseline gate-baseline.json

    # 输出每用例所属集合（便于优化报告引用）
    python3 scripts/split-gate.py results.json --split-map split-map.json

说明:
    - 切分按用例 id（无 id 时按索引）做确定性伪随机，--seed 固定则跨轮次一致。
    - 分类判定逻辑与 scripts/score-trigger-eval.py 保持一致（_classify/_compute_metrics）。
    - 回归判定容差为 0.005（F1 差绝对值 ≤0.005 视为持平，不触发回归）。
"""

from __future__ import annotations

import argparse
import json
import random
import sys
from pathlib import Path

REGRESSION_TOLERANCE = 0.005


def _load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _load_excel(path: Path) -> dict:
    """读取 Excel 触发评测结果；无 openpyxl 时报错提示。"""
    try:
        import openpyxl
    except ImportError:
        print("[error] 读取 Excel 需要 openpyxl：python3 -m pip install --break-system-packages openpyxl")
        sys.exit(1)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
    rows = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 1):
        if row[0] is None:
            continue
        d = dict(zip(headers, row))
        rounds = []
        for col in ("1轮执行结果", "2轮执行结果", "3轮执行结果"):
            rounds.append(str(d.get(col, "")).strip())
        rows.append({"id": idx, "type": str(d.get("类型", "")).strip(),
                     "expected": str(d.get("预期结果", "")).strip(), "rounds": rounds})
    return {"results": rows}


def _classify(type_val: str, expected: str, actual: str):
    """单轮分类：返回 TP/FP/FN/TN 或 None（跳过追问/无法判定）。与 score-trigger-eval.py 同步。"""
    type_val = type_val.strip()
    actual = actual.strip()

    if type_val == "正例":
        if actual == "触发":
            return "TP"
        if actual == "不触发":
            return "FN"
        return None
    if type_val in ("负例", "近似命中", "安全测试"):
        if actual == "不触发":
            return "TN"
        if actual == "触发":
            return "FP"
        return None
    if type_val == "边界情况":
        if "追问" in expected:
            if actual == "追问":
                return "TN"
            if actual == "触发":
                return "FP"
            if actual == "不触发":
                return "FN" if "应触发" in expected else "TN"
        else:
            if actual == "不触发":
                return "TN"
            if actual == "触发":
                return "FP"
            if actual == "追问":
                return "TN"
    return None


def _compute_metrics(counts: dict) -> dict:
    tp = counts.get("TP", 0)
    fp = counts.get("FP", 0)
    fn = counts.get("FN", 0)
    tn = counts.get("TN", 0)
    precision = tp / (tp + fp) if (tp + fp) else 0.0
    recall = tp / (tp + fn) if (tp + fn) else 0.0
    f1 = 2 * precision * recall / (precision + recall) if (precision + recall) else 0.0
    return {"TP": tp, "FP": fp, "FN": fn, "TN": tn,
            "precision": round(precision, 3), "recall": round(recall, 3), "f1": round(f1, 3)}


def case_key(case: dict, index: int) -> str:
    """用例稳定 id：优先取 id 字段，否则用索引。"""
    raw = case.get("id")
    if raw is None:
        return str(index)
    return str(raw)


def split_cases(results: list[dict], seed: int, holdout_ratio: float) -> tuple[list[dict], list[dict], set, set]:
    """按用例 id 确定性切分为 guide / holdout 两集（同 seed 结果一致）。

    返回 (guide_cases, holdout_cases, guide_ids, holdout_ids)。
    """
    items = [(case_key(c, i), c) for i, c in enumerate(results)]
    items.sort(key=lambda x: x[0])  # 先排序保证确定性
    rng = random.Random(seed)
    rng.shuffle(items)
    cut = max(1, int(round(len(items) * holdout_ratio)))
    holdout = [c for _, c in items[:cut]]
    guide = [c for _, c in items[cut:]]
    guide_ids = {k for k, _ in items[cut:]}
    holdout_ids = {k for k, _ in items[:cut]}
    return guide, holdout, guide_ids, holdout_ids


def metrics_for(cases: list[dict]) -> dict:
    """计算给定用例集 3 轮平均 P/R/F1。"""
    per_round = []
    for r in range(3):
        counts = {"TP": 0, "FP": 0, "FN": 0, "TN": 0}
        for item in cases:
            rounds = item.get("rounds", [])
            if r >= len(rounds):
                continue
            cls = _classify(item.get("type", ""), item.get("expected", ""), rounds[r])
            if cls:
                counts[cls] += 1
        per_round.append(_compute_metrics(counts))
    avg = {k: round(sum(m[k] for m in per_round) / len(per_round), 3)
           for k in ("precision", "recall", "f1")}
    avg["rounds"] = per_round
    return avg


def gate_verdict(current: float, baseline: float | None) -> str:
    """对比留出集 F1：improved / held / regressed。基线缺失时无法判定。"""
    if baseline is None:
        return "no-baseline"
    delta = current - baseline
    if delta > REGRESSION_TOLERANCE:
        return "improved"
    if delta < -REGRESSION_TOLERANCE:
        return "regressed"
    return "held"


def main() -> None:
    parser = argparse.ArgumentParser(description="触发评测留出验证门（guide/holdout split + 回归判定）")
    parser.add_argument("input", type=Path, help="触发评测结果 results.json 或 results.xlsx")
    parser.add_argument("--seed", type=int, default=42, help="切分随机种子（默认 42，固定则跨轮一致）")
    parser.add_argument("--holdout-ratio", type=float, default=0.2,
                        help="留出集比例（默认 0.2，即 50 条用例留出 10 条；A 卷按复杂度 25/50/100 条）")
    parser.add_argument("--baseline", type=Path, default=None,
                        help="修复前基线 JSON（gate-baseline.json），输出留出集回归判定")
    parser.add_argument("--baseline-out", type=Path, default=None,
                        help="将当前留出集指标写入基线 JSON（修复前调用一次）")
    parser.add_argument("--split-map", type=Path, default=None,
                        help="输出每用例所属集合的 JSON（guide/holdout）")
    args = parser.parse_args()

    data = _load_excel(args.input) if args.input.suffix == ".xlsx" else _load_json(args.input)
    results = data.get("results", [])
    if not results:
        print("[error] 无评测结果")
        sys.exit(1)

    guide, holdout, guide_ids, holdout_ids = split_cases(results, args.seed, args.holdout_ratio)
    guide_m = metrics_for(guide)
    hold_m = metrics_for(holdout)

    name = data.get("skill_name", "unknown")
    print(f"== 留出验证门: {name} ==")
    print(f"  用例总数 {len(results)}：引导集 {len(guide)} / 留出集 {len(holdout)}（seed={args.seed}, holdout={args.holdout_ratio:.0%}）")
    print(f"  引导集 3轮平均: P={guide_m['precision']} R={guide_m['recall']} F1={guide_m['f1']}")
    print(f"  留出集 3轮平均: P={hold_m['precision']} R={hold_m['recall']} F1={hold_m['f1']}")

    if args.baseline_out:
        baseline = {"skill_name": name, "seed": args.seed,
                    "holdout_ratio": args.holdout_ratio,
                    "holdout_f1": hold_m["f1"], "holdout_metrics": hold_m,
                    "guide_f1": guide_m["f1"], "guide_metrics": guide_m}
        args.baseline_out.write_text(
            json.dumps(baseline, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"\n  基线已写入: {args.baseline_out}（holdout F1={hold_m['f1']}）")

    if args.baseline:
        try:
            base = _load_json(args.baseline)
        except (json.JSONDecodeError, OSError) as exc:
            print(f"[error] 基线读取失败: {exc}")
            sys.exit(1)
        base_f1 = base.get("holdout_f1")
        verdict = gate_verdict(hold_m["f1"], base_f1)
        print(f"\n  基线 holdout F1={base_f1} → 当前 {hold_m['f1']}（Δ={hold_m['f1'] - base_f1:+.3f}）")
        if verdict == "regressed":
            print("  ❌ 留出集回归：本轮回合被验证门拒绝（gate_no_regression 不通过）")
            sys.exit(2)
        if verdict == "improved":
            print("  ✅ 留出集提升：验证门通过（严格提升）")
        else:
            print("  ✅ 留出集持平：验证门通过（无回归）")

    if args.split_map:
        mapping = [{"id": case_key(c, i),
                    "split": "holdout" if case_key(c, i) in holdout_ids else "guide",
                    "type": c.get("type", ""), "expected": c.get("expected", "")}
                   for i, c in enumerate(results)]
        args.split_map.write_text(json.dumps(mapping, ensure_ascii=False, indent=2),
                                  encoding="utf-8")
        print(f"\n  用例归属已写入: {args.split_map}")


if __name__ == "__main__":
    main()
