#!/usr/bin/env python3
"""入库后核验（只读）：三格表↔库逐条全等 + 非治理字段零漂移 + 与备份逐字段对照。"""
import json, os, sys, collections, hashlib
from openpyxl import load_workbook

KB = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
BK = os.path.join(KB, 'kb/l3/_backups/rule-taxonomy-20261017')
XLSX = os.path.expanduser('~/Desktop/规则清单-P0-P3-按等级-20261016.xlsx')
FIELDS = ('taxonomy_category', 'taxonomy_subcategory', 'p_level')


def rows_of(p):
    return [json.loads(l) for l in open(p, encoding='utf-8') if l.strip()]


def main():
    ok = fail = 0
    def ck(n, c, d=''):
        nonlocal ok, fail
        if c: ok += 1; print(f'  ✅ {n} {d}')
        else: fail += 1; print(f'  ❌ {n} {d}')

    wb = load_workbook(XLSX, data_only=True)
    ws = wb['全部规则 (P0-P3)']
    hdr = [ws.cell(row=5, column=c).value for c in range(1, 44)]
    tb = {}
    for r in range(6, ws.max_row + 1):
        v = [ws.cell(row=r, column=c).value for c in range(1, 44)]
        if all(x is None for x in v): continue
        d = dict(zip(hdr, v)); tb[d['规则ID(rule_id)']] = d

    now = {}
    for rel in ('kb/l3/candidates/rules.jsonl', 'kb/l3/candidates/external_standard.jsonl'):
        for r in rows_of(os.path.join(KB, rel)):
            if rel.endswith('external_standard.jsonl') or r.get('article_no'):
                now[r['rule_id']] = r

    print('\n① 三格：表 ↔ 库')
    ck('条数 638', len(now) == 638, str(len(now)))
    bad = [rid for rid, r in now.items()
           if (r.get('taxonomy_category'), r.get('taxonomy_subcategory'), r.get('p_level'))
           != (tb[rid]['所属类别'], tb[rid]['所属子类'], tb[rid]['优先级'])]
    ck('三格逐条全等', not bad, f'不符 {len(bad)}')
    dist = collections.Counter(r['p_level'] for r in now.values())
    ck('优先级分布 63/243/213/119',
       dict(dist) == {'P0': 63, 'P1': 243, 'P2': 213, 'P3': 119}, str(dict(dist)))
    ck('类别分布 11 类', len({r['taxonomy_category'] for r in now.values()}) == 11)
    ck('子类 58 个', len({r['taxonomy_subcategory'] for r in now.values()}) == 58)

    print('\n② 非治理字段零漂移（与入库前备份逐条逐字段对照）')
    for rel in ('rules.jsonl', 'external_standard.jsonl'):
        old = rows_of(os.path.join(BK, rel))
        new = rows_of(os.path.join(KB, 'kb/l3/candidates/' + rel))
        ck(f'{rel} 行数不变', len(old) == len(new), f'{len(old)} → {len(new)}')
        drift = []
        for a, b in zip(old, new):
            if a['rule_id'] != b['rule_id']:
                drift.append((a['rule_id'], '<行序>', '', '')); continue
            for k in a:
                if k in FIELDS: continue
                if a[k] != b.get(k): drift.append((a['rule_id'], k, a[k], b.get(k)))
            for k in set(b) - set(a):
                if k not in FIELDS and k != 'taxonomy_basis':
                    drift.append((a['rule_id'], '新增字段 ' + k, '', ''))
        ck(f'{rel} 非治理字段零漂移', not drift, f'漂移 {len(drift)} 处：{drift[:3]}')
        ck(f'{rel} 每条都有 taxonomy_basis',
           all('taxonomy_basis' in b for b in new), '')

    print('\n③ 关键不变量')
    p0pat = {json.loads(l)['rule_id'] for l in open(os.path.join(KB, 'kb/l5/p0_patterns.jsonl'), encoding='utf-8') if l.strip()}
    ck('快路 17 条全在 P0', all(now[r]['p_level'] == 'P0' for r in p0pat), '')
    ck('p0_patterns 未被改动（17 条）', len(p0pat) == 17)
    ck('risk_level 未被改写（仍 4 级 + unknown）',
       set(r['risk_level'] for r in now.values()) <= {'Critical', 'High', 'Medium', 'Low', 'unknown'},
       str(sorted({r['risk_level'] for r in now.values()})))
    print(f'\n合计：通过 {ok} ｜ 不通过 {fail}')
    return 0 if fail == 0 else 1


if __name__ == '__main__':
    sys.exit(main())
