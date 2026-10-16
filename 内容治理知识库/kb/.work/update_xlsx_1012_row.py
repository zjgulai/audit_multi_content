#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 10/12–10/16 行的「实际进度」按 **现跑探针** 重写一遍。

★ 为什么要走脚本、且**不覆盖原件**：
  写之前实测 `lsof` 显示 **`wpsoffice` 正持有该文件的写句柄**（`fd 100w`）＋ 有锁文件
  `~/Desktop/.~内容治理知识库搭建任务详情.xlsx` ⇒ 我若就地覆盖，**用户在 WPS 里一按保存
  就会把我的更新整份盖回去**（静默丢）。⇒ 本脚本写到**新件**，原件不动，另存一份备份。

用法：python3 kb/.work/update_xlsx_1012_row.py [--apply]
      不带 --apply 只预演（打印将要写入的两个单元格）。
"""
from __future__ import annotations

import argparse
import json
import os
import shutil
import sys
import time
from pathlib import Path

import openpyxl

DESK = Path.home() / "Desktop"
SRC = DESK / "内容治理知识库搭建任务详情.xlsx"
STAMP = time.strftime("%Y%m%d-%H%M")
OUT = DESK / ("内容治理知识库搭建任务详情-更新%s.xlsx" % STAMP)
BAK = DESK / ("内容治理知识库搭建任务详情-备份%s.xlsx" % STAMP)
ROW = 5                      # ★ 10/12–10/16 是**第 5 行**（A5；工单与结项件里写的「A6 行」是错的）

F5 = "大部分提前完成（周期未到）"

G5 = """周期未到（10/12 为本周第 1 天）· 已完成 17 项 / 部分 5 项 / 未做 2 项 / 已撤下 1 项（合计 25）
※ 下列每个数均为 2026-10-12 12:0x 现跑探针取值、不是抄旧值；出处逐条写在括号里
1. ✅ 谁负责、什么时候生效：1018/1018（rules.jsonl 行口径）
2. ⏸ 已撤下（本地自加项、PRD 无对应条款；编号保留以免下游引用错位）
3. ✅ 入库脱敏（水印 / 手机号 / 邮箱 / 小孩姓名）
4. ⚠️ 文字链路跑通：出结论、能给依据 —— ★计数口径已更正：原「累计 580 单」是数 kb/l7/runs/index.jsonl 的行数得来的；实测 source=review_machine 共 1110 行，**按素材去重只有 254 条**（重跑 856 行，单条素材最多重跑 65 次）。归因：引擎每跑一单记账，而回归跑批一次 100 单 ⇒ 度量被测量行为本身抬高。已修（写入侧隔离 KB_RUN_RECORD_SKIP）＋口径落成 kb/l8/review_run_count.py
5. ❌ 外部标准抽成规则：rules.jsonl 里 external_standard = 0 条；根因（L1 定长字符切）已归因，改动单-外部标准抽取口径-20261010.md 状态「待批」（未批不能动）
6. ✅ 风险等级逐条重判（按规则条数：Medium 217 / Low 123 / High 112 / Critical 56 / unknown 1）
7. ⚠️ 适用产品线 / 渠道：字段已建并回填 108 条；其余 401 条**源侧无对应列**（已逐条取证登记）⇒ 本地已尽、缺口在源侧
8. ⚠️ 模态：字段已建 509/509 但取值全是 undetermined；四条穷尽取证后确认**源里没有那一格**（已登记）⇒ 本地已尽
9. ✅ 内容 ID 已改业务编号 AD-YYYYMMDD-NNN（今日真单实测 AD-20261012-002）
10. ✅ WHO / AAP 条款号台账：kb/l3/registry/who-aap-clause-ledger.jsonl 40 条（回指 40/40；适用地区 3/40、生效年 11/40，其余如实留空）
11. ✅ 置信度分流策略件已建（status=active）★但阈值是 PRD 政策值、**未标定** ⇒ 标定依赖金标集（11/30 那行）
12. ✅ P1/P2 两层落地：review_engine.py 已有 P0/P1/P2/P3 四条显式分支（2026-10-10 接线）
13. ⚠️ 审计指标 7 项：4 measured / 3 no_data（覆盖率分母不存在、误漏报要人工金标、回流未接线）
14. ✅ 每周人工抽检 100 条：首轮清单已出 kb/l7/reports/audit-sample-ALL.csv（100 条）★人判栏空、待人工
15. ✅ 报告生成 ≤30 秒（实测 0.16 秒）
16. ✅ 权限按角色：矩阵件 9 个角色位已建 ★人名全空 ＋「专家」角色位无 routing key，待外部输入
17. ⚠️ 端到端性能：文字 5.1s / 图片 4.1s / 视频 5.8s 通过；**群响应 ≤5s 本机测不到**（含守衡侧＋模型侧）
18. ✅ P0 红线法务逐条确认：台账 11 条已就绪（kb/l5/registry/p0-legal-signoff-ledger.jsonl）★签署待法务（夹具就绪 ≠ 已确认）
19. ✅ 单条时效实测已出
20. ✅ 报告阈值可配置件已建 ★PDF 导出属飞书侧云文档功能，按 Owner 口径为备注项、不计交付
21. ✅ 多维表统计与筛选导出：bitable-stats-latest.json 在 ＋ BLOCK/REVIEW 两份 CSV
22. ❌ 可读性补齐（drive:file:download scope / 文档协作者）：未做 —— 缺口在飞书侧权限，非本地
23. ✅ 3 份 Brief 本地排除已定案（不动飞书源）
24. ✅ ⑥ APP资料库索引处置已落（link-overrides）
25. ✅ W1 三件输出：规则分层表 / 对接人表 / 知识库快照，三件都在 kb/l7/registry/
※ 本轮新增（不在原 25 条里；均为**实测缺陷 → 已修 ＋ 带判据**）:
 A. H09 词表变形缺口 ⇒ **真 P0 漏放**（素材逐字 baby-like natural sucking rhythm；改前整段 P0 快路 hit=False ⇒ 顶层 PASS，而历史报告判 BLOCK）→ 已补正则，matcher 自证含正例＋反例，L5 重建（regex 27→28、pattern_id 稳定），法务台账已同步
 B. 非权威候选的声明落地：seed-rules.json 自己写着 authoritative:false / 已被 kb/l3/candidates/rules.jsonl 取代，而**渲染侧原本零读者** ⇒ 5 条候选被计进「共发现 17 处必须修改」→ 已在摘要与报告就地标注（issue_count 契约值一字未改）
 C. 回归基线 ＋ 闸门：同源读法跑手 kb/l8/regression_run.py、闸门 kb/l8/regression_gate.py（solid 80 条基线 GREEN），并接进就绪巡检（**非阻断**）"""


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    a = ap.parse_args()

    if not SRC.exists():
        print("找不到原件：%s" % SRC)
        return 2
    wb = openpyxl.load_workbook(SRC)
    ws = wb["任务详情"]
    cur = ws.cell(row=ROW, column=7).value
    print("原件：%s（%d 字节）" % (SRC.name, SRC.stat().st_size))
    print("目标单元格：F%d / G%d（A%d = %r）" % (ROW, ROW, ROW, ws.cell(row=ROW, column=1).value))
    print("F%d：%r  →  %r" % (ROW, ws.cell(row=ROW, column=6).value, F5))
    print("G%d：现 %d 字  →  新 %d 字" % (ROW, len(str(cur or "")), len(G5)))
    # 会不会碰到图表/图片（openpyxl 保存会丢）
    for attr in ("_images", "_charts"):
        n = len(getattr(ws, attr, []) or [])
        if n:
            print("⚠️ 本表有 %s %d 个 ⇒ openpyxl 保存会丢，先停下" % (attr, n))
            return 3
    if not a.apply:
        print("\n预演（零写入）。落盘加 --apply。")
        return 0

    shutil.copy2(SRC, BAK)
    ws.cell(row=ROW, column=6).value = F5
    ws.cell(row=ROW, column=7).value = G5
    wb.save(OUT)

    # ★ 回读复核（不许只看"保存成功"）
    wb2 = openpyxl.load_workbook(OUT)
    w2 = wb2["任务详情"]
    ok_f = w2.cell(row=ROW, column=6).value == F5
    ok_g = w2.cell(row=ROW, column=7).value == G5
    print("\n✅ 已写入新件：%s（%d 字节）" % (OUT.name, OUT.stat().st_size))
    print("   备份原件：%s" % BAK.name)
    print("   回读复核：F%d %s ｜ G%d %s ｜ 行数 %d（原 16）"
          % (ROW, "一致 ✅" if ok_f else "**不一致** ❌",
             ROW, "一致 ✅" if ok_g else "**不一致** ❌", w2.max_row))
    print("   其余行未动：A%d=%r ｜ A6=%r" % (ROW, w2.cell(row=ROW, column=1).value,
                                             w2.cell(row=6, column=1).value))
    print("\n★ 原件**没动** —— 它正被 WPS 打开（lsof 实测 wpsoffice 持 fd 100w）。")
    print("  要就地覆盖：关掉 WPS 后 `cp '%s' '%s'`" % (OUT.name, SRC.name))
    return 0 if (ok_f and ok_g) else 3


if __name__ == "__main__":
    sys.exit(main())
